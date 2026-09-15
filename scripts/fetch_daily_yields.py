#!/usr/bin/env python3
"""Daily snapshot of US + UK 10-year government bond yields.

Sources (all keyless; run server-side so CORS is irrelevant):
  - USA: US Treasury daily par yield curve XML
        https://home.treasury.gov/sites/default/files/interest-rates/yield.xml
        (BC_10YEAR per business day; file carries roughly the latest week)
  - UK:  Bank of England nominal spot curve, daily zip of xlsx files
        https://www.bankofengland.co.uk/-/media/boe/files/statistics/yield-curves/glcnominalddata.zip
        (sheet "4. spot curve", 10-year column; latest file "...to present.xlsx")

Reads the existing data/daily-yields.json ({"USA": [[date, value], ...],
"GBR": [...]}, dates as "YYYY-MM-DD", ascending), merges in new
observations, dedupes by date and writes it back. Exits 0 with a clear
log line even if one source fails; exits 1 only if nothing could be
fetched at all AND there is no usable data to keep.
"""

import json
import os
import re
import sys
import tempfile
import urllib.request
import zipfile

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(REPO_ROOT, "data", "daily-yields.json")

TREASURY_URL = "https://home.treasury.gov/sites/default/files/interest-rates/yield.xml"
BOE_ZIP_URL = ("https://www.bankofengland.co.uk/-/media/boe/files/statistics/"
               "yield-curves/glcnominalddata.zip")

UA = {"User-Agent": "g7-bond-debt-dashboard daily snapshot (github actions)"}


def http_get(url, timeout=120):
    req = urllib.request.Request(url, headers=UA)
    return urllib.request.urlopen(req, timeout=timeout)


def fetch_us_treasury():
    """Return {date: value} from the Treasury XML daily curve."""
    with http_get(TREASURY_URL, timeout=60) as r:
        xml = r.read().decode("utf-8", errors="replace")
    out = {}
    for block in re.findall(r"<G_NEW_DATE>(.*?)</G_NEW_DATE>", xml, re.S):
        d = re.search(r"<NEW_DATE>(\d{2})-(\d{2})-(\d{4})</NEW_DATE>", block)
        v = re.search(r"<BC_10YEAR>([\d.]+)</BC_10YEAR>", block)
        if not d or not v:
            continue
        try:
            val = float(v.group(1))
        except ValueError:
            continue
        out[f"{d.group(3)}-{d.group(1)}-{d.group(2)}"] = round(val, 2)
    return out


def pick_uk_workbook(zf):
    """Pick the most recent UK nominal-curve xlsx from the BoE zip."""
    names = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
    if not names:
        raise RuntimeError("no xlsx files found in BoE zip")
    present = [n for n in names if "present" in n.lower()]
    if present:
        return sorted(present)[-1]
    def end_year(n):
        m = re.search(r"(\d{4})\s*to\s*(\d{4})", n)
        return int(m.group(2)) if m else 0
    return sorted(names, key=end_year)[-1]


def fetch_uk_boe():
    """Return {date: value} of the UK 10Y nominal spot rate from the BoE zip."""
    import openpyxl  # noqa: E402  (installed in the workflow)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")
    try:
        with http_get(BOE_ZIP_URL) as r, open(tmp.name, "wb") as f:
            while True:
                chunk = r.read(1024 * 1024)
                if not chunk:
                    break
                f.write(chunk)
        with zipfile.ZipFile(tmp.name) as zf:
            wb_name = pick_uk_workbook(zf)
            print(f"  using workbook: {wb_name}", flush=True)
            tmpdir = tempfile.mkdtemp()
            xlsx_path = zf.extract(wb_name, tmpdir)
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb["4. spot curve"]
        mat_row = list(ws.iter_rows(min_row=4, max_row=4, values_only=True))[0]
        ycol = next(i for i, v in enumerate(mat_row)
                    if str(v).strip() == "10")
        out = {}
        for row in ws.iter_rows(min_row=6, values_only=True):
            d, v = row[0], row[ycol]
            if d is None or v is None or v == "":
                continue
            ds = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)[:10]
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", ds):
                continue
            try:
                out[ds] = round(float(v), 2)
            except (TypeError, ValueError):
                continue
        wb.close()
        return out
    finally:
        os.unlink(tmp.name)


def load_existing():
    if os.path.exists(OUT_PATH):
        with open(OUT_PATH, encoding="utf-8") as f:
            d = json.load(f)
        return {k: v for k, v in d.items() if isinstance(v, list)}
    return {}


def merge(existing, new):
    merged = {p[0]: p[1] for p in existing}
    merged.update(new)
    return [[d, merged[d]] for d in sorted(merged)]


def main():
    existing = load_existing()
    failures = []

    new_us, new_uk = {}, {}
    try:
        new_us = fetch_us_treasury()
        print(f"US Treasury: {len(new_us)} daily obs "
              f"({min(new_us) if new_us else '-'}..{max(new_us) if new_us else '-'})",
              flush=True)
    except Exception as e:  # noqa: BLE001
        failures.append(f"US Treasury failed: {e}")
    try:
        new_uk = fetch_uk_boe()
        print(f"BoE UK: {len(new_uk)} daily obs "
              f"({min(new_uk) if new_uk else '-'}..{max(new_uk) if new_uk else '-'})",
              flush=True)
    except Exception as e:  # noqa: BLE001
        failures.append(f"BoE UK failed: {e}")

    for f in failures:
        print("WARNING:", f, file=sys.stderr, flush=True)

    if not new_us and not new_uk and not existing:
        print("ERROR: no data fetched and no existing snapshot; not writing.",
              file=sys.stderr)
        return 1

    data = {
        "USA": merge(existing.get("USA", []), new_us),
        "GBR": merge(existing.get("GBR", []), new_uk),
    }
    if not data["USA"] and not data["GBR"]:
        print("ERROR: merged data is empty; not writing.", file=sys.stderr)
        return 1

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    before = json.dumps(existing, sort_keys=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, separators=(",", ":"))
        f.write("\n")
    changed = json.dumps({k: data[k] for k in existing}, sort_keys=True) != before \
        or set(data) != set(existing)
    print(f"Wrote {OUT_PATH}: USA {len(data['USA'])} obs, "
          f"GBR {len(data['GBR'])} obs; changed={changed}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
